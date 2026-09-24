# Copyright (c) 2026 Luis Manuel Cousido Hermida. All rights reserved.
"""Text extraction for chat attachments (txt/code/csv/json/pdf/docx/xlsx)."""
import io
from pathlib import PurePath

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_CHARS_PER_FILE = 6000  # the local model has a small context window (MODEL_CONTEXT)

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg", ".heic"}
BINARY_EXTS = {".zip", ".rar", ".7z", ".exe", ".dll", ".mp3", ".mp4", ".wav", ".avi", ".mov", ".bin", ".iso"}


class UnsupportedFile(Exception):
    pass


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## Hoja: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                out.append(" | ".join("" if c is None else str(c) for c in row))
    return "\n".join(out)


def _text(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        raise UnsupportedFile("El archivo parece binario y no se puede leer como texto")
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_text(filename: str, data: bytes) -> dict:
    ext = PurePath(filename).suffix.lower()
    if ext in IMAGE_EXTS:
        raise UnsupportedFile("Las imágenes no son compatibles: el modelo local no tiene visión")
    if ext in BINARY_EXTS:
        raise UnsupportedFile(f"Formato {ext} no compatible")
    try:
        if ext == ".pdf":
            text = _pdf(data)
        elif ext == ".docx":
            text = _docx(data)
        elif ext in (".xlsx", ".xlsm"):
            text = _xlsx(data)
        elif ext in (".doc", ".xls", ".ppt", ".pptx", ".odt"):
            raise UnsupportedFile(f"Formato {ext} no compatible; guárdalo como .docx, .xlsx, .pdf o texto")
        else:
            text = _text(data)
    except UnsupportedFile:
        raise
    except Exception as exc:
        raise UnsupportedFile(f"No se pudo leer el archivo: {type(exc).__name__}") from exc
    text = text.strip()
    if not text:
        raise UnsupportedFile("No se encontró texto en el archivo (¿PDF escaneado?)")
    truncated = len(text) > MAX_CHARS_PER_FILE
    return {"name": filename, "text": text[:MAX_CHARS_PER_FILE], "chars": len(text), "truncated": truncated}
